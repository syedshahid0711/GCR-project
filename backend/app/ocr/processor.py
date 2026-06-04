"""OCR Processor — extract text from images, PDFs, DOCX, PPTX.

Uses free, open-source tools:
- pytesseract (Tesseract OCR) for images
- PyPDF2 for digital PDFs
- pdf2image + pytesseract for scanned PDFs
- python-docx for DOCX files
- python-pptx for PPTX files
"""
import os
import tempfile
import requests
from flask import current_app


def extract_text_from_url(url, mime_type=""):
    """Download a file from URL and extract text.

    Args:
        url: File URL (Google Drive, etc.)
        mime_type: MIME type hint

    Returns:
        Extracted text string
    """
    try:
        # Download file to temp location
        response = requests.get(url, timeout=30)
        if response.status_code != 200:
            return ""

        suffix = _get_suffix(mime_type, url)
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(response.content)
            tmp_path = tmp.name

        text = extract_text_from_file(tmp_path, mime_type)

        # Clean up
        os.unlink(tmp_path)
        return text

    except Exception as e:
        current_app.logger.error(f"OCR extraction from URL failed: {e}")
        return ""


def extract_text_from_file(file_path, mime_type=""):
    """Extract text from a local file.

    Args:
        file_path: Path to the file
        mime_type: MIME type hint

    Returns:
        Extracted text string
    """
    if not os.path.exists(file_path):
        return ""

    ext = os.path.splitext(file_path)[1].lower()
    mime = mime_type.lower()

    try:
        if ext in ('.pdf',) or 'pdf' in mime:
            return _extract_from_pdf(file_path)
        elif ext in ('.docx',) or 'wordprocessingml' in mime:
            return _extract_from_docx(file_path)
        elif ext in ('.pptx',) or 'presentationml' in mime:
            return _extract_from_pptx(file_path)
        elif ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.webp') or 'image' in mime:
            return _extract_from_image(file_path)
        elif ext in ('.txt', '.md', '.csv', '.py', '.java', '.cpp', '.js', '.html', '.css'):
            return _extract_from_text(file_path)
        else:
            current_app.logger.warning(f"Unsupported file type: {ext} ({mime})")
            return ""
    except Exception as e:
        current_app.logger.error(f"OCR extraction failed for {file_path}: {e}")
        return ""


def _extract_from_pdf(file_path):
    """Extract text from PDF — tries digital extraction first, falls back to OCR."""
    text = ""

    # Try digital PDF extraction first (fast)
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        current_app.logger.warning(f"PyPDF2 extraction failed: {e}")

    # If no text extracted, try OCR (scanned PDF)
    if not text.strip():
        try:
            from pdf2image import convert_from_path
            images = convert_from_path(file_path, dpi=200)
            for img in images:
                text += _ocr_image(img) + "\n"
        except Exception as e:
            current_app.logger.warning(f"PDF OCR fallback failed: {e}")

    return text.strip()


def _extract_from_docx(file_path):
    """Extract text from DOCX file."""
    from docx import Document
    doc = Document(file_path)
    text = "\n".join(para.text for para in doc.paragraphs if para.text)
    return text.strip()


def _extract_from_pptx(file_path):
    """Extract text from PPTX file."""
    from pptx import Presentation
    prs = Presentation(file_path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text.strip()


def _extract_from_image(file_path):
    """Extract text from image using Tesseract OCR."""
    from PIL import Image
    img = Image.open(file_path)
    return _ocr_image(img)


def _extract_from_text(file_path):
    """Extract text from plain text files."""
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read().strip()


def _ocr_image(image):
    """Run OCR on a PIL Image object using pytesseract."""
    try:
        import pytesseract
        text = pytesseract.image_to_string(image)
        return text.strip()
    except Exception as e:
        current_app.logger.error(f"Tesseract OCR failed: {e}")
        return ""


def detect_handwritten_from_image(file_path):
    """Heuristic: detect if an image contains handwritten text.

    Uses OCR confidence — low confidence often indicates handwritten content.

    Returns:
        dict with keys: is_handwritten (bool), confidence (float), text (str)
    """
    try:
        import pytesseract
        from PIL import Image

        img = Image.open(file_path)

        # Get detailed OCR data
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)

        confidences = [int(c) for c in data.get("conf", []) if int(c) > 0]

        if not confidences:
            return {"is_handwritten": True, "confidence": 90.0, "text": ""}

        avg_confidence = sum(confidences) / len(confidences)

        # Low OCR confidence suggests handwritten text
        is_handwritten = avg_confidence < 60

        text = pytesseract.image_to_string(img)

        return {
            "is_handwritten": is_handwritten,
            "confidence": 100 - avg_confidence,  # Invert: high = more likely handwritten
            "text": text.strip(),
        }

    except Exception as e:
        current_app.logger.error(f"Handwriting detection failed: {e}")
        return {"is_handwritten": False, "confidence": 0, "text": ""}


def _get_suffix(mime_type, url):
    """Get file extension from MIME type or URL."""
    mime_map = {
        "application/pdf": ".pdf",
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    }

    for key, ext in mime_map.items():
        if key in mime_type:
            return ext

    # Try from URL
    url_ext = os.path.splitext(url.split('?')[0])[1]
    return url_ext if url_ext else ".tmp"
